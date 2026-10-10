from datetime import datetime, date, timedelta
from flask import render_template, request, redirect, url_for, flash, Response
from flask_login import login_required, current_user
from sqlalchemy import func
from app.modules.reports import reports_bp
from app.models.sale import Sale, SaleItem
from app.models.product import Product
from app.models.category import Category
from app.models.expense import Expense
from app.models.customer import Customer
from app.models.supplier import Supplier
from app.models.purchase import Purchase
from app.extensions import db


def admin_required(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not current_user.is_admin:
            flash("التقارير متاحة للمدير فقط.", "danger")
            return redirect(url_for("dashboard.index"))
        return f(*args, **kwargs)
    return decorated


def _parse_date(s, fallback):
    try:
        return datetime.strptime(s, "%Y-%m-%d") if s else fallback
    except ValueError:
        return fallback


# ── Index ─────────────────────────────────────────────────────────────────────
@reports_bp.route("/")
@login_required
@admin_required
def index():
    return render_template("reports/index.html")


# ── Sales Report ──────────────────────────────────────────────────────────────
@reports_bp.route("/sales")
@login_required
@admin_required
def sales():
    report_type = request.args.get("type", "daily")
    date_str = request.args.get("date", date.today().strftime("%Y-%m-%d"))
    month_str = request.args.get("month", date.today().strftime("%Y-%m"))

    if report_type == "daily":
        try:
            sel_date = datetime.strptime(date_str, "%Y-%m-%d").date()
        except ValueError:
            sel_date = date.today()

        day_sales = Sale.query.filter(
            func.date(Sale.created_at) == sel_date,
            Sale.status != "cancelled",
        ).all()

        total_revenue = sum(s.total for s in day_sales)
        total_discount = sum(s.discount_amount for s in day_sales)
        total_profit = sum(s.profit for s in day_sales)
        cash_sales = sum(s.total for s in day_sales if s.payment_method == "cash")
        credit_sales = sum(s.total for s in day_sales if s.payment_method == "credit")
        transfer_sales = sum(s.total for s in day_sales if s.payment_method == "transfer")

        return render_template("reports/sales.html",
            report_type=report_type,
            sel_date=sel_date,
            month_str=month_str,
            day_sales=day_sales,
            total_revenue=total_revenue,
            total_discount=total_discount,
            total_profit=total_profit,
            cash_sales=cash_sales,
            credit_sales=credit_sales,
            transfer_sales=transfer_sales,
        )

    elif report_type == "monthly":
        try:
            sel_month = datetime.strptime(month_str + "-01", "%Y-%m-%d")
        except ValueError:
            sel_month = date.today().replace(day=1)

        # Days in month
        if sel_month.month == 12:
            next_month = sel_month.replace(year=sel_month.year + 1, month=1)
        else:
            next_month = sel_month.replace(month=sel_month.month + 1)

        monthly_data = db.session.query(
            func.date(Sale.created_at).label("day"),
            func.count(Sale.id).label("count"),
            func.sum(Sale.total).label("revenue"),
            func.sum(Sale.subtotal - Sale.total).label("discount"),
        ).filter(
            Sale.created_at >= sel_month,
            Sale.created_at < next_month,
            Sale.status != "cancelled",
        ).group_by(func.date(Sale.created_at)).order_by("day").all()

        month_total = sum(r.revenue or 0 for r in monthly_data)
        month_count = sum(r.count for r in monthly_data)

        return render_template("reports/sales.html",
            report_type=report_type,
            date_str=date_str,
            month_str=month_str,
            sel_month=sel_month,
            monthly_data=monthly_data,
            month_total=month_total,
            month_count=month_count,
        )

    elif report_type == "cashier":
        from app.models.user import User
        cashier_data = db.session.query(
            Sale.user_id,
            func.count(Sale.id).label("count"),
            func.sum(Sale.total).label("revenue"),
        ).filter(Sale.status != "cancelled").group_by(Sale.user_id).all()

        cashier_rows = []
        for row in cashier_data:
            user = User.query.get(row.user_id)
            cashier_rows.append({
                "name": user.full_name if user else "—",
                "count": row.count,
                "revenue": row.revenue or 0,
            })

        return render_template("reports/sales.html",
            report_type=report_type,
            date_str=date_str,
            month_str=month_str,
            cashier_rows=cashier_rows,
        )

    return redirect(url_for("reports.sales"))


# ── Profit Report ─────────────────────────────────────────────────────────────
@reports_bp.route("/profit")
@login_required
@admin_required
def profit():
    today = date.today()
    first_of_month = today.replace(day=1)

    date_from_str = request.args.get("date_from", first_of_month.strftime("%Y-%m-%d"))
    date_to_str = request.args.get("date_to", today.strftime("%Y-%m-%d"))

    date_from = _parse_date(date_from_str, datetime.combine(first_of_month, datetime.min.time()))
    date_to = _parse_date(date_to_str, datetime.combine(today, datetime.max.time()))
    date_to = date_to.replace(hour=23, minute=59, second=59)

    # Sales in period
    period_sales = Sale.query.filter(
        Sale.created_at >= date_from,
        Sale.created_at <= date_to,
        Sale.status != "cancelled",
    ).all()

    total_revenue = sum(s.total for s in period_sales)
    total_cogs = sum(
        item.purchase_price * item.quantity
        for s in period_sales
        for item in s.items
    )
    gross_profit = total_revenue - total_cogs

    # Expenses in period
    total_expenses = db.session.query(
        func.coalesce(func.sum(Expense.amount), 0)
    ).filter(
        Expense.expense_date >= date_from,
        Expense.expense_date <= date_to,
    ).scalar()

    net_profit = gross_profit - float(total_expenses)

    # Profit by product
    product_profit = db.session.query(
        Product.name,
        func.sum(SaleItem.quantity).label("qty"),
        func.sum(SaleItem.total).label("revenue"),
        func.sum(SaleItem.purchase_price * SaleItem.quantity).label("cogs"),
    ).join(SaleItem, SaleItem.product_id == Product.id)\
     .join(Sale, Sale.id == SaleItem.sale_id)\
     .filter(
        Sale.created_at >= date_from,
        Sale.created_at <= date_to,
        Sale.status != "cancelled",
     ).group_by(Product.id)\
      .order_by(func.sum(SaleItem.total - SaleItem.purchase_price * SaleItem.quantity).desc())\
      .limit(20).all()

    # Profit by category
    cat_profit = db.session.query(
        Category.name,
        func.sum(SaleItem.total).label("revenue"),
        func.sum(SaleItem.purchase_price * SaleItem.quantity).label("cogs"),
    ).join(Product, Product.id == SaleItem.product_id)\
     .join(Category, Category.id == Product.category_id)\
     .join(Sale, Sale.id == SaleItem.sale_id)\
     .filter(
        Sale.created_at >= date_from,
        Sale.created_at <= date_to,
        Sale.status != "cancelled",
     ).group_by(Category.id)\
      .order_by(func.sum(SaleItem.total - SaleItem.purchase_price * SaleItem.quantity).desc())\
      .all()

    return render_template("reports/profit.html",
        date_from_str=date_from_str,
        date_to_str=date_to_str,
        total_revenue=total_revenue,
        total_cogs=total_cogs,
        gross_profit=gross_profit,
        total_expenses=float(total_expenses),
        net_profit=net_profit,
        product_profit=product_profit,
        cat_profit=cat_profit,
        sale_count=len(period_sales),
    )


# ── Products Report ───────────────────────────────────────────────────────────
@reports_bp.route("/products")
@login_required
@admin_required
def products():
    today = date.today()
    first_of_month = today.replace(day=1)

    date_from_str = request.args.get("date_from", first_of_month.strftime("%Y-%m-%d"))
    date_to_str = request.args.get("date_to", today.strftime("%Y-%m-%d"))
    idle_days = int(request.args.get("idle_days", 30))

    date_from = _parse_date(date_from_str, datetime.combine(first_of_month, datetime.min.time()))
    date_to = _parse_date(date_to_str, datetime.combine(today, datetime.max.time())).replace(hour=23, minute=59, second=59)

    # Top selling by qty
    top_qty = db.session.query(
        Product.name,
        Product.unit,
        func.sum(SaleItem.quantity).label("total_qty"),
        func.sum(SaleItem.total).label("total_rev"),
    ).join(SaleItem, SaleItem.product_id == Product.id)\
     .join(Sale, Sale.id == SaleItem.sale_id)\
     .filter(
        Sale.created_at >= date_from,
        Sale.created_at <= date_to,
        Sale.status != "cancelled",
     ).group_by(Product.id)\
      .order_by(func.sum(SaleItem.quantity).desc())\
      .limit(15).all()

    # Top selling by revenue
    top_rev = db.session.query(
        Product.name,
        func.sum(SaleItem.total).label("total_rev"),
    ).join(SaleItem, SaleItem.product_id == Product.id)\
     .join(Sale, Sale.id == SaleItem.sale_id)\
     .filter(
        Sale.created_at >= date_from,
        Sale.created_at <= date_to,
        Sale.status != "cancelled",
     ).group_by(Product.id)\
      .order_by(func.sum(SaleItem.total).desc())\
      .limit(15).all()

    # Idle products
    idle_cutoff = datetime.combine(today - timedelta(days=idle_days), datetime.min.time())
    sold_ids = db.session.query(SaleItem.product_id).join(Sale)\
        .filter(Sale.created_at >= idle_cutoff, Sale.status != "cancelled")\
        .distinct().subquery()

    idle_products = Product.query.filter(
        Product.is_active == True,
        Product.quantity > 0,
        ~Product.id.in_(db.session.query(sold_ids)),
    ).order_by(Product.quantity.desc()).all()

    # Low stock
    low_stock = Product.query.filter(
        Product.is_active == True,
        Product.quantity <= Product.low_stock_alert,
    ).order_by(Product.quantity).all()

    return render_template("reports/products.html",
        date_from_str=date_from_str,
        date_to_str=date_to_str,
        idle_days=idle_days,
        top_qty=top_qty,
        top_rev=top_rev,
        idle_products=idle_products,
        low_stock=low_stock,
    )


# ── Accounts Report ───────────────────────────────────────────────────────────
@reports_bp.route("/accounts")
@login_required
@admin_required
def accounts():
    # Customer debts
    customers_with_debt = [
        c for c in Customer.query.filter_by(is_active=True).all()
        if c.balance > 0
    ]
    total_customer_debt = sum(c.balance for c in customers_with_debt)

    # Supplier debts
    suppliers_with_debt = [
        s for s in Supplier.query.filter_by(is_active=True).all()
        if s.balance > 0
    ]
    total_supplier_debt = sum(s.balance for s in suppliers_with_debt)

    # Stock value
    stock_at_cost = db.session.query(
        func.coalesce(func.sum(Product.quantity * Product.purchase_price), 0)
    ).filter_by(is_active=True).scalar()

    stock_at_price = db.session.query(
        func.coalesce(func.sum(Product.quantity * Product.sale_price), 0)
    ).filter_by(is_active=True).scalar()

    stock_potential_profit = float(stock_at_price) - float(stock_at_cost)

    # Products count
    total_products = Product.query.filter_by(is_active=True).count()
    out_of_stock = Product.query.filter(
        Product.is_active == True, Product.quantity <= 0
    ).count()

    return render_template("reports/accounts.html",
        customers_with_debt=customers_with_debt,
        total_customer_debt=total_customer_debt,
        suppliers_with_debt=suppliers_with_debt,
        total_supplier_debt=total_supplier_debt,
        stock_at_cost=float(stock_at_cost),
        stock_at_price=float(stock_at_price),
        stock_potential_profit=stock_potential_profit,
        total_products=total_products,
        out_of_stock=out_of_stock,
    )


# ── Excel Export: Sales ───────────────────────────────────────────────────────
@reports_bp.route("/sales/export/excel")
@login_required
@admin_required
def sales_export_excel():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    import io

    date_from_str = request.args.get("date_from", date.today().strftime("%Y-%m-%d"))
    date_to_str = request.args.get("date_to", date.today().strftime("%Y-%m-%d"))

    date_from = _parse_date(date_from_str, datetime.combine(date.today(), datetime.min.time()))
    date_to = _parse_date(date_to_str, datetime.combine(date.today(), datetime.max.time())).replace(hour=23,minute=59,second=59)

    sales = Sale.query.filter(
        Sale.created_at >= date_from,
        Sale.created_at <= date_to,
        Sale.status != "cancelled",
    ).order_by(Sale.created_at).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "المبيعات"
    ws.sheet_view.rightToLeft = True

    hfill = PatternFill("solid", fgColor="1e3a5f")
    hfont = Font(bold=True, color="FFFFFF")

    headers = ["#", "رقم الفاتورة", "العميل", "الكاشير", "المجموع", "الخصم", "الصافي", "طريقة الدفع", "التاريخ"]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = hfont
        c.fill = hfill
        c.alignment = Alignment(horizontal="center")

    for row, s in enumerate(sales, 2):
        ws.cell(row=row, column=1, value=row-1)
        ws.cell(row=row, column=2, value=s.invoice_number)
        ws.cell(row=row, column=3, value=s.customer.name if s.customer else "نقدي")
        ws.cell(row=row, column=4, value=s.cashier.full_name)
        ws.cell(row=row, column=5, value=s.subtotal)
        ws.cell(row=row, column=6, value=s.discount_amount)
        ws.cell(row=row, column=7, value=s.total)
        ws.cell(row=row, column=8, value={"cash":"نقدي","credit":"آجل","transfer":"تحويل"}.get(s.payment_method,""))
        ws.cell(row=row, column=9, value=s.created_at.strftime("%Y-%m-%d %H:%M"))

    total_row = len(sales) + 2
    ws.cell(row=total_row, column=6, value="الإجمالي").font = Font(bold=True)
    ws.cell(row=total_row, column=7, value=sum(s.total for s in sales)).font = Font(bold=True)

    for col in ["A","B","C","D","E","F","G","H","I"]:
        ws.column_dimensions[col].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=sales-report.xlsx"},
    )